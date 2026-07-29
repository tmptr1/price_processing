from PySide6.QtCore import QThread, Signal
from PySide6.QtWidgets import QFileDialog
import time
import traceback
import os
import datetime
from sqlalchemy import func
# from sqlalchemy.orm import sessionmaker
# from sqlalchemy.exc import OperationalError, UnboundExecutionError
# from sqlalchemy import select, delete, insert, and_
import pandas as pd
import colors
import openpyxl
import smtplib
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.header import Header
from email import encoders
# from CatalogUpdate import get_tables_skip_row_dict

import setting
# engine = setting.get_engine()
# session = sessionmaker(engine)
settings_data = setting.get_vars()

LOG_ID = 5


class OrderSenderClass(QThread):
    StartSignal = Signal(bool)
    default_settings = True
    need_to_send = True
    morning = False
    evening = False
    file_path = ''
    sheet_name = ''
    table_name = ''
    dir_path = ''
    skip_rows = 0
    errors = 0
    tables_skip_row_dict = dict()  # название таблицы: [название листа, skip rows]

    def __init__(self, log=None, parent=None):
        self.log = log
        QThread.__init__(self, parent)


    def run(self):
        # global session, engine
        self.errors = 0

        try:
            extra_info = {'утро': self.morning, 'evening': self.evening}
            extra_info = [k if extra_info[k] else False for k in extra_info.keys()]
            type_list = list(filter(lambda x: x, extra_info))
            extra_info = f"({','.join(type_list)})"
            if type_list == []:
                self.log.add(LOG_ID, f"Необходимо выбрать время заказа: утро, evening",
                             f"Необходимо выбрать время заказа: <span style='color:{colors.orange_log_color};'>утро, evening</span>  ")
                return

            if not self.file_path:
                self.log.add(LOG_ID, f"Необходимо выбрать путь к файлу",
                             f"<span style='color:{colors.orange_log_color};'>Необходимо выбрать путь к файлу</span>  ")
                return
            if not self.dir_path:
                self.log.add(LOG_ID, f"Необходимо выбрать папку для сохранения",
                             f"<span style='color:{colors.orange_log_color};'>Необходимо выбрать папку для сохранения</span>  ")
                return

            self.log.add(LOG_ID, f"Старт {extra_info}", f"<span style='color:{colors.green_log_color};'>Старт</span> {extra_info}  ")

            if not self.sheet_name:
                self.sheet_name = self.tables_skip_row_dict[self.table_name][0]

            if not self.default_settings:
                skip_rows = self.tables_skip_row_dict.get(self.table_name)[1]
            else:
                skip_rows = 0

            df = pd.read_excel(self.file_path, sheet_name=self.sheet_name, skiprows=skip_rows)
            cols = {"Заказ", "Поставщик", "Ключ1 в заказ", "Артикул в заказ", "Производитель в заказ", "Наименование в заказ",
                    "Количество", "Цена в заказ", "Сумма", "Изменение количества. Новое, если отличается",
                    "Новая цена", "Адрес"}
            # "Кратность в заказ", "Примечание в заказ"
            if not cols.issubset(set(df.columns.tolist())):
                self.log.add(LOG_ID, f"Нет необходимых столбцов", f"<span style='color:{colors.orange_log_color};'>Нет необходимых столбцов</span>  ")
                return

            df = df[df['Заказ'].isin(type_list)]
            if df.empty:
                self.log.add(LOG_ID, f"Заказы не найдены",
                             f"<span style='color:{colors.orange_log_color};'>Заказы не найдены</span>  ")
                return

            self.log.add(LOG_ID, f"Папка для прайсов: {self.dir_path}")
            self.log.add(LOG_ID, f"----------------")

            cols_for_csv = ["Ключ1 в заказ", "Артикул в заказ", "Производитель в заказ", "Наименование в заказ",
                    "Количество", "Цена в заказ", "Сумма", "Изменение количества. Новое, если отличается", "Новая цена"]
            self.now_date = datetime.datetime.now().date().strftime('%d.%m.%y')
            for day_part in type_list:
                suppliers = df[(df['Заказ']==day_part) & (df['Поставщик'].notna())]['Поставщик'].unique()
                for supplier in suppliers:
                    new_df = df[df['Поставщик']==supplier]
                    self.file_name = fr"{supplier} {self.now_date} {day_part}.xlsx"  # csv
                    self.log.add(LOG_ID, f"{self.file_name} - формирование...")
                    self.send_to = new_df.iloc[0]['Адрес']
                    new_df = new_df[cols_for_csv]
                    # new_df.to_csv(fr"{self.dir_path}/{self.file_name}", sep=';', decimal=',', encoding="windows-1251", index=False, errors='ignore')
                    new_df.to_excel(fr"{self.dir_path}/{self.file_name}", index=False)
                    self.log.add(LOG_ID, f"{self.file_name} сформирован", f"{self.file_name} <span style='color:{colors.green_log_color};'>сформирован</span>  ")
                    if '@' in f"{self.send_to}" and self.need_to_send:
                        if not self.send_email(day_part):
                            self.errors += 1
                        break
                    else:
                        self.log.add(LOG_ID, f"Не указана почта для отправки / Не указана отправка",
                                     f"<span style='color:{colors.orange_log_color};'>Не указана почта для отправки / Не указана отправка</span>  ")

                    self.log.add(LOG_ID, f"")

            self.log.add(LOG_ID, f"Заказы отправлены")
            if self.errors:
                self.log.add(LOG_ID, f"Не отправлено файлов: {self.errors}",
                             f"<span style='color:{colors.orange_log_color};'>Не отправлено файлов:</span> {self.errors}  ")

        # except (OperationalError, UnboundExecutionError) as db_ex:
        #     self.log.add(LOG_ID, f"Повторное подключение к БД ...", f"<span style='color:{colors.orange_log_color};"
        #                                                         f"font-weight:bold;'>Повторное подключение к БД ...</span>  ")
            # try:
            #     engine = setting.get_engine()
            #     session = sessionmaker(engine)
            # except:
            #     pass
        except Exception as ex:
            ex_text = traceback.format_exc()
            self.log.error(LOG_ID, "ERROR", ex_text)


    def send_email(self, day_part):
        try:
            day_part = {'утро': 'утренний', 'evening': 'вечерний'}[day_part]
            for send_to_ in [self.send_to, 'ytopttorg@mail.ru']:
                msg = MIMEMultipart()
                msg["Subject"] = Header(f"Заказ от ИП Шевелько ({self.now_date} {day_part})")
                msg["From"] = settings_data['mail_orders_login']
                msg["To"] = send_to_

                # s = smtplib.SMTP("smtp.yandex.ru", 587, timeout=100)
                with smtplib.SMTP_SSL("smtp.mail.ru", 465, timeout=100) as s:
                    # s.starttls()
                    s.login(settings_data['mail_orders_login'], settings_data['mail_orders_imap_password'])

                    file_path = fr"{self.dir_path}/{self.file_name}"
                    with open(file_path, 'rb') as f:
                        file = MIMEBase('application', 'vnd.ms-excel')
                        file.set_payload(f.read())

                    encoders.encode_base64(file)
                    file.add_header('Content-Disposition', 'attachment', filename=os.path.basename(file_path))
                    msg.attach(file)

                    s.sendmail(msg["From"], send_to_, msg.as_string())


            self.log.add(LOG_ID, f"Отправлено на {self.send_to}",
                         f"<span style='color:{colors.green_log_color};'>Отправлено</span> на {self.send_to}  ")
            return True

        except smtplib.SMTPDataError as smtp_ex:
            ex_text = traceback.format_exc()
            self.log.error(LOG_ID, f"Подозрение на спам ({smtp_ex.smtp_code}) ERROR:", ex_text)

        except Exception as send_ex:
            ex_text = traceback.format_exc()
            self.log.error(LOG_ID, "SEND ERROR", ex_text)

        return False

    def select_file(self):
        try:
            self.custom_file_path = QFileDialog.getOpenFileName(filter='Excel File (*.xlsx *.xls)')[0]

        except Exception as ex:
            ex_text = traceback.format_exc()
            self.log.error(LOG_ID, "Ошибка при выборе файла", ex_text)

    def select_path_to_save(self):
        try:
            self.path_to_save = QFileDialog.getExistingDirectory()
        except Exception as ex:
            ex_text = traceback.format_exc()
            self.log.error(LOG_ID, "Ошибка при выборе файла", ex_text)



class GetSheets(QThread):
    FinishSignal = Signal(dict)
    def __init__(self, log, custom_file_path, parent=None):
        self.custom_file_path = custom_file_path
        self.log = log
        QThread.__init__(self, parent)

    def run(self):
        try:
            self.log.add(LOG_ID, f"Загрузка таблиц...")
            self.tables_skip_row_dict = get_tables_skip_row_dict(self.custom_file_path)
            self.FinishSignal.emit(self.tables_skip_row_dict)
            self.log.add(LOG_ID, f"Таблицы загружены")
        except Exception as ex:
            ex_text = traceback.format_exc()
            self.log.error(LOG_ID, "Ошибка при загрузке таблицы", ex_text)



def get_tables_skip_row_dict(path_to_file):
    workbook = openpyxl.load_workbook(filename=path_to_file)
    tables_skip_rows_dict = dict()
    for ws in workbook.worksheets:
        for t in ws.tables:
            # print(t)
            table = ws.tables[t]
            # print(table)
            # print(ws.title)
            # print(table.ref)
            skip_r = int(''.join(filter(str.isdigit, (str(table.ref).split(':')[0])))) - 1
            tables_skip_rows_dict[t] = [ws.title, skip_r]

    return tables_skip_rows_dict